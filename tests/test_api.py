"""Run from demo root: python -m unittest discover -s tests -p test_*.py -v."""

import concurrent.futures
import importlib
import io
import json
import os
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
from datetime import date
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE))
TEMP = tempfile.TemporaryDirectory(prefix="medical-rag-tests-")
os.environ["MEDICAL_RAG_DATA_DIR"] = TEMP.name
os.environ["MEDICAL_RAG_AS_OF_DATE"] = "2026-08-31"
from server.app import app
from server import config, db, text2sql
from server.rag import ENGINE, semantic_split
from server.documents import DOCUMENTS
from server.executor import execute_sql
from server.tasks import TaskService
from fastapi.testclient import TestClient

api_module = importlib.import_module("server.app")
CLIENT = None


def setUpModule():
    global CLIENT
    CLIENT = TestClient(app)
    CLIENT.__enter__()


def tearDownModule():
    CLIENT.__exit__(None, None, None)
    TEMP.cleanup()


def ask(content, **kw):
    return CLIENT.post("/api/chat/send", json={"content": content, **kw})


def oracle(sql):
    c = db.biz_conn()
    try:
        return [list(r) for r in c.execute(sql).fetchall()]
    finally:
        c.close()


# Each expected query is written independently of the planner under test.
SQL_CASES = [
    (
        "2025年销售额",
        "SELECT ROUND(SUM(amount),2) FROM fact_sales WHERE date_id BETWEEN '2025-01-01' AND '2025-12-31'",
    ),
    (
        "2025年第一季度销售额",
        "SELECT ROUND(SUM(amount),2) FROM fact_sales WHERE date_id BETWEEN '2025-01-01' AND '2025-03-31'",
    ),
    (
        "2025年第2季度销售额",
        "SELECT ROUND(SUM(amount),2) FROM fact_sales WHERE date_id BETWEEN '2025-04-01' AND '2025-06-30'",
    ),
    (
        "2025年Q3销售额",
        "SELECT ROUND(SUM(amount),2) FROM fact_sales WHERE date_id BETWEEN '2025-07-01' AND '2025-09-30'",
    ),
    (
        "2025年第四季度销售额",
        "SELECT ROUND(SUM(amount),2) FROM fact_sales WHERE date_id BETWEEN '2025-10-01' AND '2025-12-31'",
    ),
    (
        "2025年上半年销售额",
        "SELECT ROUND(SUM(amount),2) FROM fact_sales WHERE date_id BETWEEN '2025-01-01' AND '2025-06-30'",
    ),
    (
        "2025年下半年销售额",
        "SELECT ROUND(SUM(amount),2) FROM fact_sales WHERE date_id BETWEEN '2025-07-01' AND '2025-12-31'",
    ),
    (
        "2025年10月销售额",
        "SELECT ROUND(SUM(amount),2) FROM fact_sales WHERE date_id BETWEEN '2025-10-01' AND '2025-10-31'",
    ),
    (
        "去年销售额",
        "SELECT ROUND(SUM(amount),2) FROM fact_sales WHERE date_id BETWEEN '2025-01-01' AND '2025-12-31'",
    ),
    (
        "本月销售额",
        "SELECT ROUND(SUM(amount),2) FROM fact_sales WHERE date_id BETWEEN '2026-08-01' AND '2026-08-31'",
    ),
    (
        "上个月销售额",
        "SELECT ROUND(SUM(amount),2) FROM fact_sales WHERE date_id BETWEEN '2026-07-01' AND '2026-07-31'",
    ),
    (
        "近3个月销售额",
        "SELECT ROUND(SUM(amount),2) FROM fact_sales WHERE date_id BETWEEN '2026-05-01' AND '2026-07-31'",
    ),
    (
        "2025年销量",
        "SELECT SUM(quantity) FROM fact_sales WHERE date_id BETWEEN '2025-01-01' AND '2025-12-31'",
    ),
    (
        "2025年成本",
        "SELECT ROUND(SUM(cost_amount),2) FROM fact_sales WHERE date_id BETWEEN '2025-01-01' AND '2025-12-31'",
    ),
    (
        "2025年毛利",
        "SELECT ROUND(SUM(amount-cost_amount),2) FROM fact_sales WHERE date_id BETWEEN '2025-01-01' AND '2025-12-31'",
    ),
    (
        "2025年毛利率",
        "SELECT ROUND((1-SUM(cost_amount)/SUM(amount))*100,2) FROM fact_sales WHERE date_id BETWEEN '2025-01-01' AND '2025-12-31'",
    ),
    ("平均销售单价", "SELECT ROUND(SUM(amount)/SUM(quantity),2) FROM fact_sales"),
    (
        "阿莫西林销售额",
        "SELECT ROUND(SUM(f.amount),2) FROM fact_sales f JOIN dim_drug d ON d.drug_id=f.drug_id WHERE d.drug_name LIKE '阿莫西林%'",
    ),
    (
        "上海销售额",
        "SELECT ROUND(SUM(f.amount),2) FROM fact_sales f JOIN dim_store s ON s.store_id=f.store_id WHERE s.city='上海'",
    ),
    (
        "上海和杭州销售额",
        "SELECT ROUND(SUM(f.amount),2) FROM fact_sales f JOIN dim_store s ON s.store_id=f.store_id WHERE s.city IN ('上海','杭州')",
    ),
    (
        "华东区域销售额",
        "SELECT s.region,ROUND(SUM(f.amount),2) FROM fact_sales f JOIN dim_store s ON s.store_id=f.store_id WHERE s.region='华东' GROUP BY s.region",
    ),
    (
        "中成药销售额",
        "SELECT ROUND(SUM(f.amount),2) FROM fact_sales f JOIN dim_drug d ON d.drug_id=f.drug_id WHERE d.category='中成药'",
    ),
    (
        "直营销售额",
        "SELECT s.store_type,ROUND(SUM(f.amount),2) FROM fact_sales f JOIN dim_store s ON s.store_id=f.store_id WHERE s.store_type='直营' GROUP BY s.store_type",
    ),
    ("采购额", "SELECT ROUND(SUM(amount),2) FROM fact_purchase"),
    ("采购数量", "SELECT SUM(quantity) FROM fact_purchase"),
    (
        "上海供应商采购额",
        "SELECT s.supplier_name,ROUND(SUM(f.amount),2) FROM fact_purchase f JOIN dim_supplier s ON s.supplier_id=f.supplier_id WHERE s.city='上海' GROUP BY s.supplier_name ORDER BY 2 DESC",
    ),
    (
        "当前库存金额",
        "SELECT ROUND(SUM(stock_amount),2) FROM fact_inventory WHERE date_id='2026-08-28'",
    ),
    (
        "当前库存数量",
        "SELECT SUM(stock_qty) FROM fact_inventory WHERE date_id='2026-08-28'",
    ),
    (
        "2025年12月库存金额",
        "SELECT ROUND(SUM(stock_amount),2) FROM fact_inventory WHERE date_id='2025-12-28'",
    ),
    (
        "近效期库存记录数",
        "SELECT COUNT(*) FROM fact_inventory WHERE date_id='2026-08-28' AND expiry_date BETWEEN '2026-08-31' AND date('2026-08-31','+180 days')",
    ),
    (
        "已过期库存记录数",
        "SELECT COUNT(*) FROM fact_inventory WHERE date_id='2026-08-28' AND expiry_date<'2026-08-31'",
    ),
    (
        "销量前10的药品",
        "SELECT d.drug_name,SUM(f.quantity) FROM fact_sales f JOIN dim_drug d ON d.drug_id=f.drug_id GROUP BY d.drug_name ORDER BY 2 DESC LIMIT 10",
    ),
    (
        "库存金额最高的5个药品",
        "SELECT d.drug_name,ROUND(SUM(f.stock_amount),2) FROM fact_inventory f JOIN dim_drug d ON d.drug_id=f.drug_id WHERE f.date_id='2026-08-28' GROUP BY d.drug_name ORDER BY 2 DESC LIMIT 5",
    ),
    (
        "各月份销售额趋势",
        "SELECT substr(date_id,1,7),ROUND(SUM(amount),2) FROM fact_sales GROUP BY substr(date_id,1,7) ORDER BY 1",
    ),
    (
        "各品类销售额",
        "SELECT d.category,ROUND(SUM(f.amount),2) FROM fact_sales f JOIN dim_drug d ON d.drug_id=f.drug_id GROUP BY d.category ORDER BY 2 DESC",
    ),
    (
        "销量最低的3个药品",
        "SELECT d.drug_name,SUM(f.quantity) FROM fact_sales f JOIN dim_drug d ON d.drug_id=f.drug_id GROUP BY d.drug_name ORDER BY 2 ASC LIMIT 3",
    ),
]


class SqlGoldenTests(unittest.TestCase):
    pass


def sql_case(q, expected_sql):
    def test(self):
        response = ask(q)
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()["data"]["aiMessage"]["extraData"]
        self.assertEqual(result["error"], "", result)
        self.assertEqual(result["data"], oracle(expected_sql), q)

    return test


for i, (q, sql) in enumerate(SQL_CASES):
    setattr(SqlGoldenTests, f"test_sql_{i + 1:02d}", sql_case(q, sql))

RAG_CASES = [
    ("处方药销售有什么要求？", ["处方", "药师"]),
    ("药品储存的温湿度要求是什么？", ["10 至 30", "百分之三十五"]),
    ("近效期药品怎么分级预警？", ["十二个月", "六个月", "三个月"]),
    ("供应商评估为 C 级会怎么处理？", ["C 级", "暂停采购"]),
    ("药品召回分为几级，时限分别是多少？", ["二十四小时", "四十八小时", "七十二小时"]),
    ("冷链药品交接有什么要求？", ["交接时间", "三十分钟"]),
    ("冷链到货验收需要检查什么？", ["运输温度记录", "拒收"]),
    ("处方审核发现问题怎么处理？", ["拒绝调配", "处方医生"]),
    ("药师不在岗时销售有什么限制？", ["暂停销售"]),
    ("库存近效期一级预警怎么处理？", ["停止向门店配送"]),
    ("供应商连续两个周期评估为C级？", ["移出", "两年"]),
    ("药品召回结束后报告包括什么？", ["召回原因", "数量"]),
    ("过期药品销毁需要记录什么？", ["影像记录", "双人签字"]),
    ("销售凭证需要哪些字段？", ["药品名称", "数量"]),
    ("量子纠缠和黑洞", []),
    ("编写太空旅行诗歌", []),
]


class RagGoldenTests(unittest.TestCase):
    pass


def rag_case(q, keys):
    def test(self):
        response = ask(q, module="QA")
        self.assertEqual(response.status_code, 200, response.text)
        msg = response.json()["data"]["aiMessage"]
        for key in keys:
            self.assertIn(key, msg["content"])
        if not keys:
            self.assertEqual(msg["referDoc"], [])
        for ref in msg["referDoc"]:
            preview = CLIENT.get("/api/chat/pdf/preview/" + str(ref["docId"])).text
            self.assertIn(ref["snippet"], preview)
        self.assertFalse(msg["extraData"]["trace"][0]["detail"]["rerankEnabled"])

    return test


for i, (q, keys) in enumerate(RAG_CASES):
    setattr(RagGoldenTests, f"test_rag_{i + 1:02d}", rag_case(q, keys))


class RegressionTests(unittest.TestCase):
    def test_reject_unsupported_and_invalid(self):
        for q in [
            "帮我写一首诗",
            "2025年13月销售额",
            "销售额同比",
            "客单价",
            "销量前99999的药品",
            "近0个月销售额",
            "采购门店金额",
            "库存批次数",
            "销售额和销量",
        ]:
            with self.subTest(q=q):
                self.assertEqual(ask(q).status_code, 400)
        for kw in [{"module": "INVALID"}, {"sessionId": -1}, {"content": "x" * 2001}]:
            self.assertEqual(
                CLIENT.post(
                    "/api/chat/send", json={"content": "销售额", **kw}
                ).status_code,
                422,
            )
        self.assertEqual(ask("销售额", dbConfigId=99999).status_code, 400)
        self.assertEqual(ask("销售额", sessionId=99999).status_code, 400)

    def test_fixed_clock_and_empty(self):
        self.assertEqual(
            text2sql.parse_time("上个月", date(2026, 10, 1)),
            ("2026-09-01", "2026-09-30"),
        )
        self.assertEqual(
            text2sql.parse_time("2025年第一季度", date(2026, 10, 1)),
            ("2025-01-01", "2025-03-31"),
        )
        result = ask("2028年销售额").json()["data"]["aiMessage"]
        self.assertNotIn("None", result["content"])
        self.assertEqual(result["extraData"]["data"], [])

    def test_self_correction_preserves_metric(self):
        for q in ["2025年销售额", "销量前10的药品", "库存金额最高的5个药品"]:
            with self.subTest(q=q):
                a = ask(q).json()["data"]["aiMessage"]["extraData"]
                b = ask(q, forceError=True).json()["data"]["aiMessage"]["extraData"]
                self.assertEqual(a["data"], b["data"])
                self.assertEqual(b["error"], "")
                self.assertTrue(any(t["step"] == "自纠错重写" for t in b["trace"]))

    def test_regenerate_keeps_versions_and_one_user(self):
        result = ask("2025年销售额").json()["data"]
        sid = result["sessionId"]
        mid = result["aiMessage"]["id"]
        r = CLIENT.post(
            "/api/chat/regenerate", json={"messageId": mid, "module": "DATA"}
        )
        self.assertEqual(r.status_code, 200, r.text)
        rows = CLIENT.get("/api/chat/messages", params={"sessionId": sid}).json()[
            "data"
        ]
        self.assertEqual([m["role"] for m in rows], ["user", "assistant"])
        self.assertEqual(rows[-1]["supersedes"], mid)
        with db.platform_db() as c:
            self.assertEqual(
                c.execute(
                    "SELECT active FROM chat_message WHERE id=?", (mid,)
                ).fetchone()[0],
                0,
            )
        self.assertEqual(
            CLIENT.post(
                "/api/chat/regenerate",
                json={"messageId": rows[-1]["id"], "module": "QA"},
            ).status_code,
            400,
        )
        with patch.object(
            api_module, "run_question", side_effect=ValueError("simulated failure")
        ):
            self.assertEqual(
                CLIENT.post(
                    "/api/chat/regenerate",
                    json={"messageId": rows[-1]["id"], "module": "DATA"},
                ).status_code,
                400,
            )
        unchanged = CLIENT.get("/api/chat/messages", params={"sessionId": sid}).json()[
            "data"
        ]
        self.assertEqual(rows, unchanged)

    def test_cancellation_does_not_publish(self):
        started = threading.Event()

        def slow(*args, **kwargs):
            started.set()
            self.assertTrue(kwargs["cancel"].wait(3))
            raise InterruptedError()

        with (
            patch.object(api_module, "run_question", side_effect=slow),
            concurrent.futures.ThreadPoolExecutor() as pool,
        ):
            future = pool.submit(ask, "销售额", taskId="cancel-regression")
            self.assertTrue(started.wait(2))
            reply = CLIENT.post(
                "/api/chat/cancel", json={"taskId": "cancel-regression"}
            )
            self.assertEqual(reply.json()["data"]["state"], "cancelled")
            self.assertEqual(future.result(timeout=4).status_code, 409)
        self.assertEqual(
            CLIENT.get("/api/chat/tasks/cancel-regression").json()["data"]["state"],
            "cancelled",
        )
        with db.platform_db() as c:
            self.assertEqual(
                c.execute(
                    "SELECT status FROM ask_log ORDER BY id DESC LIMIT 1"
                ).fetchone()[0],
                "cancelled",
            )

    def test_cancel_before_registration_and_terminal_state(self):
        service = TaskService()
        service.cancel("early")
        with self.assertRaises(InterruptedError):
            service.start("early", None)
        t = service.start("normal", None)
        self.assertEqual(service.publish(t, lambda: 42), 42)
        service.finish(t)
        self.assertEqual(service.cancel("normal")["state"], "completed")
        with self.assertRaises(ValueError):
            service.start("normal", None)

    def test_readonly_and_timeout(self):
        self.assertEqual(
            execute_sql("SELECT 'update' AS harmless")["data"], [["update"]]
        )
        for sql in [
            "DELETE FROM fact_sales",
            "SELECT 1; SELECT 2",
            "SELECT name FROM sqlite_master",
            "SELECT load_extension('anything')",
            "SELECT randomblob(1000000000)",
        ]:
            with self.subTest(sql=sql), self.assertRaises((ValueError, sqlite3.Error)):
                execute_sql(sql)
        with self.assertRaises(TimeoutError):
            execute_sql(
                "SELECT SUM(a.amount*b.amount) FROM fact_sales a CROSS JOIN fact_sales b",
                timeout=0.01,
            )
        cancelled = threading.Event()
        cancelled.set()
        with self.assertRaises(InterruptedError):
            execute_sql("SELECT 1", cancel=cancelled)
        conn = db.biz_conn()
        try:
            with self.assertRaises(sqlite3.OperationalError):
                conn.execute("DELETE FROM fact_sales")
        finally:
            conn.close()

    def test_query_export_limits_and_formula_safety(self):
        result = ask("销售明细").json()["data"]
        extra = result["aiMessage"]["extraData"]
        self.assertEqual(extra["returnedRows"], 500)
        self.assertTrue(extra["hasMore"])
        response = CLIENT.get(
            "/api/chat/export/excel", params={"messageId": result["aiMessage"]["id"]}
        )
        wb = load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.active.max_row, 501)
        with db.platform_db() as c:
            c.execute(
                "UPDATE chat_message SET extra_data=? WHERE id=?",
                (
                    json.dumps({"columns": ["text"], "data": [["=1+1"]]}),
                    result["aiMessage"]["id"],
                ),
            )
        response = CLIENT.get(
            "/api/chat/export/excel", params={"messageId": result["aiMessage"]["id"]}
        )
        wb = load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.active["A2"].data_type, "s")
        self.assertEqual(
            CLIENT.get(
                "/api/chat/export/excel", params={"messageId": 999999}
            ).status_code,
            404,
        )

    def test_upload_validation(self):
        for name, content in [
            ("../escape.txt", b"x"),
            ("C:\\escape.txt", b"x"),
            ("/absolute.txt", b"x"),
            ("a.pdf", b"%PDF"),
            ("a.txt", b"\xff"),
            ("empty.txt", b""),
        ]:
            r = CLIENT.post("/api/admin/doc/upload", files={"file": (name, content)})
            with self.subTest(name=name):
                self.assertEqual(r.status_code, 400, r.text)
        self.assertEqual(
            CLIENT.post(
                "/api/admin/doc/upload",
                files={"file": ("big.txt", b"x" * (config.MAX_UPLOAD_BYTES + 1))},
            ).status_code,
            413,
        )
        self.assertFalse((Path(TEMP.name) / "escape.txt").exists())

    def test_document_versions_and_rebuild(self):
        def upload(name, text):
            r = CLIENT.post(
                "/api/admin/doc/upload", files={"file": (name, text.encode())}
            )
            self.assertEqual(r.status_code, 200, r.text)
            return r.json()["data"]

        old = upload("zzz-regression.txt", "斑马运输：交接需要双人复核。")
        before = CLIENT.get("/api/chat/pdf/preview/" + str(old["id"])).text
        upload("aaa-regression.txt", "苹果仓库：每天盘点一次。")
        new = upload("zzz-regression.txt", "斑马运输：现在需要三人复核。")
        self.assertNotEqual(old["id"], new["id"])
        self.assertEqual(new["version"], 2)
        self.assertEqual(
            upload("zzz-regression.txt", "斑马运输：现在需要三人复核。")["id"],
            new["id"],
        )
        DOCUMENTS.initialize()
        ENGINE.build()
        self.assertEqual(
            CLIENT.get("/api/chat/pdf/preview/" + str(old["id"])).text, before
        )
        docs = CLIENT.get("/api/admin/doc/list").json()["data"]
        matching = [d for d in docs if d["doc_name"] == "zzz-regression.txt"]
        self.assertEqual(len(matching), 1)
        self.assertEqual(matching[0]["versions"], 2)
        self.assertIn("三人复核", ENGINE.answer("斑马运输交接")["content"])

    def test_index_snapshot_during_build(self):
        old = ENGINE.chunks
        with patch.object(
            DOCUMENTS, "indexed", side_effect=RuntimeError("simulated index failure")
        ):
            with self.assertRaises(RuntimeError):
                ENGINE.build()
        self.assertIs(ENGINE.chunks, old)
        self.assertTrue(ENGINE.answer("冷链交接")["referDoc"])
        self.assertTrue(all(len(c) <= 128 for c in semantic_split("甲" * 1000)))

    def test_dashboard_metrics_and_validation(self):
        detail = CLIENT.get("/api/dashboard/detail/1").json()["data"]
        for card in detail["config"]["cards"]:
            r = CLIENT.post("/api/dashboard/executeSql", json={"sql": card["sql"]})
            self.assertEqual(r.status_code, 200, r.text)
            if card["id"] == "c4":
                self.assertEqual(
                    r.json()["data"]["data"],
                    ask("近效期库存记录数").json()["data"]["aiMessage"]["extraData"][
                        "data"
                    ],
                )
        saved = CLIENT.post(
            "/api/dashboard/save",
            json={
                "name": "regression",
                "refreshMode": "off",
                "config": detail["config"],
            },
        )
        self.assertEqual(saved.status_code, 200, saved.text)
        loaded = CLIENT.get(
            "/api/dashboard/detail/" + str(saved.json()["data"]["id"])
        ).json()["data"]
        self.assertEqual(loaded["refresh_mode"], "off")
        config_copy = json.loads(json.dumps(detail["config"]))
        config_copy["cards"][0]["x"] = 11
        self.assertEqual(
            CLIENT.post(
                "/api/dashboard/save", json={"name": "bad", "config": config_copy}
            ).status_code,
            422,
        )
        self.assertEqual(
            CLIENT.post(
                "/api/dashboard/executeSql", json={"sql": "SELECT 1", "dbConfigId": 2}
            ).status_code,
            422,
        )

    def test_config_truth_and_stats(self):
        self.assertFalse(
            CLIENT.post("/api/admin/ds/test", json={"id": 2}).json()["data"]["ok"]
        )
        self.assertTrue(
            CLIENT.post("/api/admin/ds/test", json={"id": 1}).json()["data"]["ok"]
        )
        self.assertFalse(
            CLIENT.post("/api/admin/model/test", json={"endpoint": "invalid"}).json()[
                "data"
            ]["ok"]
        )
        self.assertEqual(CLIENT.post("/api/admin/ds/save", json={}).status_code, 400)
        ask("2025年销售额")
        ask("量子纠缠和黑洞", module="QA")
        ask("2025年13月销售额")
        rows = CLIENT.get("/api/admin/stats/overview").json()["data"]["recent"]
        self.assertEqual(rows[0]["status"], "rejected")
        self.assertEqual(rows[1]["status"], "no_evidence")
        self.assertIn("SELECT", rows[2]["sql_text"])

    def test_idempotent_database_migration(self):
        before = oracle("SELECT COUNT(*) FROM fact_sales")
        db.init_all()
        self.assertEqual(before, oracle("SELECT COUNT(*) FROM fact_sales"))

    def test_legacy_schema_and_upload_adoption(self):
        with tempfile.TemporaryDirectory(prefix="medical-rag-legacy-") as legacy:
            code = """
import sqlite3
from pathlib import Path
from server import config,db
from server.documents import DOCUMENTS
from server.rag import ENGINE
p=Path(config.DATA_DIR); (p/'uploads').mkdir()
(p/'uploads'/'legacy.txt').write_text('旧文档：运输交接需要签字。',encoding='utf8')
c=sqlite3.connect(config.PLATFORM_DB); c.executescript(db._platform_schema_sql())
c.execute("INSERT INTO chat_session(id,module,title,created_at) VALUES (1,'DATA','旧会话','2025-01-01')")
c.execute("INSERT INTO chat_message(session_id,role,content) VALUES (1,'user','销售额')")
c.execute("INSERT INTO chat_message(session_id,role,content) VALUES (1,'assistant','旧回答')")
c.commit();c.close()
db.init_all();DOCUMENTS.initialize();ENGINE.build()
with db.platform_db() as c:
 row=c.execute('SELECT reply_to,active,content FROM chat_message WHERE id=2').fetchone()
 assert tuple(row)==(1,1,'旧回答'),tuple(row)
doc=[d for d in DOCUMENTS.list() if d['doc_name']=='legacy.txt'][0]
assert doc['id']==6
assert '签字' in DOCUMENTS.get(6)['content']
"""
            result = subprocess.run(
                [sys.executable, "-X", "utf8", "-c", code],
                cwd=BASE,
                env=dict(os.environ, MEDICAL_RAG_DATA_DIR=legacy),
                capture_output=True,
                text=True,
                encoding="utf8",
            )
            self.assertEqual(result.returncode, 0, result.stderr)

    def test_long_document_and_concurrent_read(self):
        content = "甲" * 2400 + "\n\n斑马验收必须复核封签与运输单据。"
        with concurrent.futures.ThreadPoolExecutor() as pool:
            reads = [pool.submit(ask, "冷链交接", module="QA") for _ in range(4)]
            response = CLIENT.post(
                "/api/admin/doc/upload",
                files={"file": ("long-evidence.txt", content.encode())},
            )
            self.assertEqual(response.status_code, 200, response.text)
            for future in reads:
                r = future.result(timeout=4)
                self.assertEqual(r.status_code, 200, r.text)
                self.assertIn("三十分钟", r.json()["data"]["aiMessage"]["content"])
        answer = ENGINE.answer("斑马验收复核")
        self.assertIn("封签", answer["content"])
        self.assertNotIn("甲" * 30, answer["content"])


if __name__ == "__main__":
    unittest.main()
